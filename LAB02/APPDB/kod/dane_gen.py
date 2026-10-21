"""Generator danych syntetycznych KawaNaKlik (APPDB, Etap I).

Tworzy arkusz dane/KawaNaKlik_dane_syntetyczne.xlsx (klienci, produkty,
zamówienia, pozycje, e-maile, faktury), te same tabele w CSV oraz pliki HTML
faktur kosztowych (konwertowane potem do PDF).

Wszystkie osoby, firmy, NIP-y i numery kont są fikcyjne. E-maile klientów
są w zarezerwowanej domenie .test.

Uruchomienie: python dane_gen.py <katalog_dane>
"""

import csv
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from emaile_llm import EMAILE

random.seed(2026)  # powtarzalne dane

IMIONA_K = ["Anna", "Katarzyna", "Magdalena", "Agnieszka", "Joanna", "Ewa", "Monika", "Natalia",
            "Aleksandra", "Zofia", "Julia", "Karolina", "Paulina", "Marta", "Weronika", "Alicja"]
IMIONA_M = ["Piotr", "Tomasz", "Paweł", "Marcin", "Michał", "Krzysztof", "Jakub", "Kamil",
            "Mateusz", "Łukasz", "Adam", "Bartosz", "Szymon", "Wojciech", "Filip", "Jan"]
NAZWISKA = [("Nowak", "Nowak"), ("Wójcik", "Wójcik"), ("Kowalczyk", "Kowalczyk"),
            ("Zieliński", "Zielińska"), ("Wiśniewski", "Wiśniewska"), ("Lewandowski", "Lewandowska"),
            ("Kamiński", "Kamińska"), ("Dąbrowski", "Dąbrowska"), ("Mazur", "Mazur"),
            ("Krawczyk", "Krawczyk"), ("Piotrowski", "Piotrowska"), ("Grabowski", "Grabowska"),
            ("Pawlak", "Pawlak"), ("Michalski", "Michalska"), ("Sowa", "Sowa"), ("Lis", "Lis"),
            ("Jabłoński", "Jabłońska"), ("Wróbel", "Wróbel"), ("Sikora", "Sikora"), ("Baran", "Baran")]
MIASTA = [("Warszawa", 0.22), ("Kraków", 0.14), ("Wrocław", 0.09), ("Poznań", 0.08), ("Gdańsk", 0.07),
          ("Łódź", 0.07), ("Katowice", 0.05), ("Lublin", 0.05), ("Szczecin", 0.04), ("Bydgoszcz", 0.04),
          ("Rzeszów", 0.04), ("Białystok", 0.03), ("Toruń", 0.03), ("Olsztyn", 0.03), ("Kielce", 0.02)]

PRODUKTY = [
    # SKU, nazwa, kategoria, cena brutto, stan, waga kg, popularność
    ("KNK-ETI-250", "Etiopia Yirgacheffe 250 g", "Kawa ziarnista", 49.90, 140, 0.25, 9),
    ("KNK-ETI-1000", "Etiopia Yirgacheffe 1 kg", "Kawa ziarnista", 169.00, 35, 1.0, 3),
    ("KNK-KOL-250", "Kolumbia Supremo 250 g", "Kawa ziarnista", 42.90, 160, 0.25, 8),
    ("KNK-KOL-1000", "Kolumbia Supremo 1 kg", "Kawa ziarnista", 129.00, 60, 1.0, 5),
    ("KNK-BRA-500", "Brazylia Santos 500 g", "Kawa ziarnista", 64.50, 90, 0.5, 6),
    ("KNK-BRA-1000", "Brazylia Santos 1 kg", "Kawa ziarnista", 115.00, 40, 1.0, 4),
    ("KNK-ESP-1000", "Espresso Blend 1 kg", "Kawa ziarnista", 99.00, 12, 1.0, 7),
    ("KNK-GWA-250", "Gwatemala Antigua 250 g", "Kawa ziarnista", 46.90, 80, 0.25, 4),
    ("KNK-KEN-250", "Kenia AA 250 g", "Kawa ziarnista", 54.90, 45, 0.25, 3),
    ("KNK-RWA-250", "Rwanda Nyamasheke 250 g", "Kawa ziarnista", 52.90, 30, 0.25, 2),
    ("KNK-SUM-250", "Sumatra Mandheling 250 g", "Kawa ziarnista", 48.90, 50, 0.25, 2),
    ("KNK-GWA-250M", "Gwatemala Antigua mielona 250 g", "Kawa mielona", 44.90, 110, 0.25, 6),
    ("KNK-KOL-250M", "Kolumbia Supremo mielona 250 g", "Kawa mielona", 42.90, 120, 0.25, 6),
    ("KNK-MEX-250M", "Bezkofeinowa Meksyk mielona 250 g", "Kawa mielona", 46.90, 8, 0.25, 2),
    ("KNK-ESP-250M", "Espresso Blend mielona 250 g", "Kawa mielona", 34.90, 70, 0.25, 5),
    ("KNK-V60", "Drip V60 ceramiczny", "Akcesoria", 89.00, 25, 0.4, 2),
    ("KNK-V60-PL", "Drip V60 plastikowy", "Akcesoria", 39.00, 40, 0.15, 2),
    ("KNK-CHE-6", "Chemex 6 filiżanek", "Akcesoria", 219.00, 10, 0.9, 1),
    ("KNK-AERO", "Aeropress", "Akcesoria", 179.00, 18, 0.5, 2),
    ("KNK-MLY", "Młynek ręczny stalowy", "Akcesoria", 159.00, 15, 0.7, 2),
    ("KNK-MLY-EL", "Młynek elektryczny żarnowy", "Akcesoria", 449.00, 6, 2.5, 1),
    ("KNK-FIL-100", "Filtry papierowe 02 – 100 szt.", "Akcesoria", 19.90, 300, 0.1, 5),
    ("KNK-FIL-CHE", "Filtry do Chemexa – 100 szt.", "Akcesoria", 49.00, 60, 0.2, 1),
    ("KNK-WAG", "Waga baristyczna z timerem", "Akcesoria", 99.00, 20, 0.4, 1),
    ("KNK-CZA", "Czajnik z wąską wylewką 0,8 l", "Akcesoria", 149.00, 14, 0.8, 1),
    ("KNK-KUB", "Kubek ceramiczny KawaNaKlik", "Akcesoria", 39.00, 80, 0.4, 2),
    ("KNK-DEG-4", "Zestaw degustacyjny 4 × 100 g", "Zestawy prezentowe", 79.00, 45, 0.5, 4),
    ("KNK-POR", "Zestaw „Poranek” (kawa + kubek)", "Zestawy prezentowe", 119.00, 0, 0.8, 2),
    ("KNK-PRE", "Zestaw „Pour-over Start”", "Zestawy prezentowe", 199.00, 12, 1.2, 1),
    ("KNK-SUB-1", "Subskrypcja: 1 kawa / miesiąc", "Subskrypcje", 45.00, 999, 0.25, 2),
]

STATUSY = [("zrealizowane", 0.78), ("w realizacji", 0.09), ("oczekuje na płatność", 0.04),
           ("anulowane", 0.03), ("zwrócone", 0.02), ("wstrzymane", 0.04)]
PLATNOSCI = [("BLIK", 0.38), ("karta", 0.24), ("przelew tradycyjny", 0.26), ("pobranie", 0.12)]
WYSYLKI = [("Paczkomat", 0.58, 12.99), ("Kurier", 0.42, 14.99)]
BLEDY = ["literówka w adresie przy przepisywaniu", "etykieta nadana na zły adres",
         "pominięty produkt przy kompletacji", "zły produkt (mielona zamiast ziarnistej)",
         "nie wysłano numeru przesyłki", "przelew nieprzypisany > 2 dni"]


def pick(weighted):
    r, acc = random.random(), 0.0
    for item in weighted:
        acc += item[1]
        if r <= acc:
            return item
    return weighted[-1]


def klienci(n=1200):
    out = []
    for i in range(1, n + 1):
        kobieta = random.random() < 0.55
        imie = random.choice(IMIONA_K if kobieta else IMIONA_M)
        naz = random.choice(NAZWISKA)[1 if kobieta else 0]
        miasto = pick(MIASTA)[0]
        login = (imie + "." + naz).lower()
        for a, b in zip("ąćęłńóśźż", "acelnoszz"):
            login = login.replace(a, b)
        out.append(dict(
            klient_id=f"K{i:04d}", imie=imie, nazwisko=naz, email=f"{login}{i}@kawanaklik.test",
            telefon=f"500 {random.randint(100, 999)} {i:03d}"[:11], miasto=miasto,
            kod_pocztowy=f"{random.randint(0, 99):02d}-{random.randint(0, 999):03d}",
            ulica=f"ul. Testowa {random.randint(1, 120)}",
            typ="firma" if random.random() < 0.08 else "osoba prywatna",
            zgoda_marketing=random.random() < 0.41,
            data_rejestracji=(date(2024, 1, 1) + timedelta(days=random.randint(0, 640))).isoformat(),
        ))
    return out


def zamowienia(kl, dni=14, start=date(2026, 10, 1)):
    zam, poz = [], []
    waga = [p[6] for p in PRODUKTY]
    # Firmy i stali klienci zamawiają częściej
    aktywnosc = [random.lognormvariate(0, 0.7) * (3 if k["typ"] == "firma" else 1) for k in kl]
    nr = 10000
    for d in range(dni):
        dzien = start + timedelta(days=d)
        baza = {0: 135, 1: 118, 2: 112, 3: 110, 4: 104, 5: 82, 6: 79}[dzien.weekday()]
        for _ in range(int(random.gauss(baza, 9))):
            nr += 1
            k = random.choices(kl, weights=aktywnosc)[0]
            ts = datetime.combine(dzien, datetime.min.time()) + timedelta(
                minutes=int(random.triangular(6 * 60, 23 * 60 + 59, 20 * 60)))
            pozycje = random.choices(PRODUKTY, weights=waga, k=random.choices([1, 2, 3, 4], [0.45, 0.33, 0.15, 0.07])[0])
            suma = 0.0
            for p in {p[0]: p for p in pozycje}.values():
                ile = random.choices([1, 2, 3], [0.8, 0.16, 0.04])[0]
                suma += ile * p[3]
                poz.append(dict(nr_zamowienia=f"KNK-{nr}", sku=p[0], produkt=p[1], ilosc=ile,
                                cena_jedn=p[3], wartosc=round(ile * p[3], 2)))
            wys = pick(WYSYLKI)
            koszt = 0.0 if suma >= 199 else wys[2]
            status = pick(STATUSY)[0]
            if (date(2026, 10, 14) - dzien).days < 2 and status == "zrealizowane":
                status = "w realizacji"
            blad = random.choice(BLEDY) if random.random() < 0.024 else ""
            zam.append(dict(
                nr_zamowienia=f"KNK-{nr}", data=ts.strftime("%Y-%m-%d %H:%M"), klient_id=k["klient_id"],
                miasto=k["miasto"], status=status, platnosc=pick(PLATNOSCI)[0], wysylka=wys[0],
                liczba_pozycji=len({p[0] for p in pozycje}), wartosc_produktow=round(suma, 2),
                koszt_wysylki=koszt, razem=round(suma + koszt, 2), blad_reczny=blad,
            ))
    return zam, poz


def emaile(kl, zam):
    zam_klienta = {}
    for z in zam:
        zam_klienta.setdefault(z["klient_id"], []).append(z)
    kandydaci = [k for k in kl if k["klient_id"] in zam_klienta]
    out = []
    for i, e in enumerate(EMAILE, 1):
        k = random.choice(kandydaci)
        z = random.choice(zam_klienta[k["klient_id"]])
        ts = datetime.strptime(z["data"], "%Y-%m-%d %H:%M") + timedelta(hours=random.randint(2, 96))
        spam = e["kategoria"] == "spam"
        b2b = e["kategoria"] == "oferta_b2b"
        nadawca = "kontakt@seo-rocket.test" if spam else k["email"]
        imie = "" if spam else f'{k["imie"]} {k["nazwisko"]}'
        uzyto_zam = "{zam}" in e["tresc"] or "{zam}" in e["temat"]
        out.append(dict(
            email_id=f"M{i:03d}", data=min(ts, datetime(2026, 10, 18, 22, 0)).strftime("%Y-%m-%d %H:%M"),
            nadawca=nadawca, klient_id="" if spam or b2b else k["klient_id"],
            nr_zamowienia=z["nr_zamowienia"] if uzyto_zam else "",
            temat=e["temat"].replace("{zam}", z["nr_zamowienia"]),
            tresc=e["tresc"].replace("{zam}", z["nr_zamowienia"]).replace("{imie}", imie).strip(),
            kategoria_wzorcowa=e["kategoria"], priorytet_wzorcowy=e["priorytet"], ton=e["ton"],
        ))
    return sorted(out, key=lambda r: r["data"])


DOSTAWCY = [
    dict(nazwa="Palarnia Ziarno i Ogień Sp. z o.o.", adres="ul. Fikcyjna 7, 30-000 Kraków", nip="999-001-01-01",
         pozycje=[("Kawa palona Etiopia Yirgacheffe – worek 10 kg", 6, 520.00),
                  ("Kawa palona Kolumbia Supremo – worek 10 kg", 8, 430.00),
                  ("Kawa palona Espresso Blend – worek 10 kg", 5, 360.00)], termin=14),
    dict(nazwa="EkoPak Opakowania S.A.", adres="ul. Testowa 44, 90-000 Łódź", nip="999-002-02-02",
         pozycje=[("Torebka 250 g z zaworem, nadruk – 1000 szt.", 3, 690.00),
                  ("Torebka 1 kg z zaworem, nadruk – 500 szt.", 1, 540.00),
                  ("Karton wysyłkowy 20×15×10 – 500 szt.", 2, 410.00)], termin=21),
    dict(nazwa="KurierExpress Polska Sp. z o.o.", adres="ul. Przykładowa 3, 02-000 Warszawa", nip="999-003-03-03",
         pozycje=[("Usługi kurierskie – wrzesień 2026 (612 przesyłek)", 1, 7340.40),
                  ("Dopłata paliwowa", 1, 512.80)], termin=14),
    dict(nazwa="HostNet Sp. z o.o.", adres="ul. Serwerowa 1, 50-000 Wrocław", nip="999-004-04-04",
         pozycje=[("Hosting WordPress Business – 12 mies.", 1, 1188.00),
                  ("Certyfikat SSL Wildcard – 12 mies.", 1, 290.00)], termin=7),
    dict(nazwa="DrukArt Studio Graficzne", adres="ul. Kolorowa 12, 80-000 Gdańsk", nip="999-005-05-05",
         pozycje=[("Ulotki A6 do paczek – 5000 szt.", 1, 450.00),
                  ("Naklejki z logo – 2000 szt.", 1, 380.00)], termin=14),
    dict(nazwa="Młynki Pro Hurt Sp. j.", adres="ul. Żarnowa 9, 60-000 Poznań", nip="999-006-06-06",
         pozycje=[("Młynek ręczny stalowy – hurt", 20, 92.00),
                  ("Drip V60 ceramiczny – hurt", 24, 48.00),
                  ("Waga baristyczna z timerem – hurt", 10, 55.00)], termin=30,
         uwaga="Na zamówieniu było 20 szt. drippera V60, na fakturze 24 szt. – faktura do wyjaśnienia."),
]


def zl(x):
    return f"{x:,.2f}".replace(",", " ").replace(".", ",")


def faktury(katalog: Path):
    out = []
    katalog.mkdir(parents=True, exist_ok=True)
    for i, d in enumerate(DOSTAWCY, 1):
        nr = f"FV/{i:02d}/10/2026"
        wyst = date(2026, 10, 1) + timedelta(days=random.randint(0, 12))
        netto = round(sum(q * c for _, q, c in d["pozycje"]), 2)
        vat = round(netto * 0.23, 2)
        wiersze = "".join(
            f"<tr><td>{j}</td><td>{n}</td><td class=r>{q}</td><td class=r>{zl(c)}</td>"
            f"<td class=r>23%</td><td class=r>{zl(q * c)}</td></tr>"
            for j, (n, q, c) in enumerate(d["pozycje"], 1))
        html = f"""<!doctype html><html lang=pl><meta charset=utf-8><style>
@page{{size:A4;margin:16mm}}body{{font-family:Segoe UI,Arial,sans-serif;font-size:10pt;color:#222}}
h1{{font-size:18pt;margin:0}}.top{{display:flex;justify-content:space-between;margin-bottom:8mm}}
.box{{width:48%;border:1px solid #bbb;padding:3mm;min-height:24mm}}.flex{{display:flex;justify-content:space-between}}
table{{width:100%;border-collapse:collapse;margin-top:6mm}}th,td{{border:1px solid #bbb;padding:1.5mm 2mm}}
th{{background:#eee;text-align:left}}.r{{text-align:right}}.sum td{{font-weight:700}}
.wm{{position:fixed;top:45%;left:5%;font-size:44pt;color:rgba(200,0,0,.12);transform:rotate(-25deg);font-weight:700}}
.note{{margin-top:6mm;font-size:8.5pt;color:#666}}</style>
<div class=wm>DOKUMENT TESTOWY</div>
<div class=top><div><h1>Faktura VAT</h1><div>nr <b>{nr}</b></div></div>
<div class=r>Data wystawienia: {wyst:%d.%m.%Y}<br>Data sprzedaży: {wyst:%d.%m.%Y}<br>
Termin płatności: {(wyst + timedelta(days=d['termin'])):%d.%m.%Y}</div></div>
<div class=flex><div class=box><b>Sprzedawca</b><br>{d['nazwa']}<br>{d['adres']}<br>NIP: {d['nip']}<br>
Rachunek: PL00 0000 0000 0000 0000 0000 00{i:02d}</div>
<div class=box><b>Nabywca</b><br>KawaNaKlik Sp. z o.o.<br>ul. Testowa 1, 31-000 Kraków<br>NIP: 999-000-00-00</div></div>
<table><tr><th>Lp.</th><th>Nazwa towaru / usługi</th><th class=r>Ilość</th><th class=r>Cena netto</th>
<th class=r>VAT</th><th class=r>Wartość netto</th></tr>{wiersze}
<tr class=sum><td colspan=5 class=r>Razem netto</td><td class=r>{zl(netto)}</td></tr>
<tr class=sum><td colspan=5 class=r>VAT 23%</td><td class=r>{zl(vat)}</td></tr>
<tr class=sum><td colspan=5 class=r>Do zapłaty (brutto) PLN</td><td class=r>{zl(netto + vat)}</td></tr></table>
<p>Forma płatności: przelew</p>
<p class=note>Faktura wygenerowana na potrzeby laboratorium APPDB. Wszystkie dane (firmy, NIP, rachunki) są fikcyjne.</p>
</html>"""
        plik = f"faktura_{i:02d}_{d['nazwa'].split()[0].lower().replace('ł', 'l')}.html"
        (katalog / plik).write_text(html, encoding="utf-8")
        out.append(dict(nr_faktury=nr, dostawca=d["nazwa"], nip=d["nip"], data_wystawienia=wyst.isoformat(),
                        termin_platnosci=(wyst + timedelta(days=d["termin"])).isoformat(),
                        netto=netto, vat=vat, brutto=round(netto + vat, 2),
                        plik_pdf=plik.replace(".html", ".pdf"), uwagi=d.get("uwaga", "")))
    return out


def zapisz(katalog: Path, tabele: dict):
    wb = Workbook()
    wb.remove(wb.active)
    naglowek = PatternFill("solid", fgColor="6B3E26")
    for nazwa, wiersze in tabele.items():
        ws = wb.create_sheet(nazwa)
        kolumny = list(wiersze[0].keys())
        ws.append(kolumny)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = naglowek
        for w in wiersze:
            ws.append([w[k] for k in kolumny])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, k in enumerate(kolumny, 1):
            szer = min(60, max(len(k), *(len(str(w[k])) for w in wiersze[:200])) + 2)
            ws.column_dimensions[get_column_letter(i)].width = szer
            if k == "tresc":
                for c in ws[get_column_letter(i)][1:]:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        with open(katalog / f"{nazwa}.csv", "w", newline="", encoding="utf-8-sig") as f:
            wr = csv.DictWriter(f, fieldnames=kolumny, delimiter=";")
            wr.writeheader()
            wr.writerows(wiersze)
    wb.save(katalog / "KawaNaKlik_dane_syntetyczne.xlsx")


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else "dane")
    out.mkdir(parents=True, exist_ok=True)
    kl = klienci()
    zam, poz = zamowienia(kl)
    prod = [dict(sku=p[0], nazwa=p[1], kategoria=p[2], cena_brutto=p[3], stan_magazynowy=p[4], waga_kg=p[5])
            for p in PRODUKTY]
    mail = emaile(kl, zam)
    fv = faktury(out / "faktury")
    zapisz(out, dict(klienci=kl, produkty=prod, zamowienia=zam, pozycje_zamowien=poz, emaile=mail, faktury=fv))
    print(f"klienci={len(kl)} produkty={len(prod)} zamowienia={len(zam)} pozycje={len(poz)} "
          f"emaile={len(mail)} faktury={len(fv)} bledy={sum(1 for z in zam if z['blad_reczny'])}")
